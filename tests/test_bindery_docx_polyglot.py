# -*- coding: utf-8 -*-
"""
Tests for DOCX Polyglot Parallel Table Grid (DocxPolyglotBuilder)
验证 Word 多语言对照版导出：段落级水平锁定表格、表头徽章、无边框装帧与整体渲染。
"""

import os
import tempfile
import docx
from core.bindery.docx_polyglot_builder import DocxPolyglotBuilder
from core.adapters.egress.ebook.docx import DocxBookAdapter


def test_docx_polyglot_split_markdown_blocks():
    """验证 Markdown 文本能精准切分为语义块：标题、代码块与段落"""
    md = """# 一级标题

这是第一段正文描述。

```python
def hello():
    return "world"
```

这是第二段正文分析。
"""
    blocks = DocxPolyglotBuilder.split_markdown_blocks(md)
    assert len(blocks) == 4
    assert blocks[0].startswith("# 一级标题")
    assert "第一段正文" in blocks[1]
    assert blocks[2].startswith("```python")
    assert "第二段正文" in blocks[3]


def test_docx_polyglot_render_chapter_table():
    """验证多语言对照表格生成：2 栏、表头带 [ZH] 与 [EN] 徽章、内容严格水平对齐"""
    doc = docx.Document()
    ch = {
        "title": "数字出版概览",
        "chapter_by_lang": {
            "zh": {
                "raw_body": "# 概览\n\n数字出版推动内容自主。\n\n关键特性在于多端自愈。"
            },
            "en": {
                "raw_body": "# Overview\n\nDigital publishing empowers content sovereignty.\n\nKey feature is multi-terminal healing."
            }
        }
    }
    DocxPolyglotBuilder.render_polyglot_chapter(doc, ch, ["zh", "en"])

    assert len(doc.tables) == 1
    tbl = doc.tables[0]
    # 表头行 + 2 个正文内容块行 (首行重复 H1 已剥离，第1段 + 第2段) = 3 行
    assert len(tbl.rows) == 3
    assert len(tbl.columns) == 2

    # 验证表头单元格
    hdr_zh = tbl.cell(0, 0).text
    hdr_en = tbl.cell(0, 1).text
    assert "[ZH]" in hdr_zh
    assert "[EN]" in hdr_en

    # 验证水平对照行 1 (第1段直接水平对齐)
    assert "推动内容自主" in tbl.cell(1, 0).text
    assert "content sovereignty" in tbl.cell(1, 1).text

    # 验证水平对照行 2 (第2段)
    assert "多端自愈" in tbl.cell(2, 0).text
    assert "multi-terminal healing" in tbl.cell(2, 1).text


def test_docx_adapter_full_polyglot_export():
    """验证通过 DocxBookAdapter 完整导出多语言对照 Word 文档"""
    adapter = DocxBookAdapter()
    manuscript_tree = [
        {
            "id": "ch_1",
            "title": "绪论与宗旨",
            "languages": ["zh", "en"],
            "chapter_by_lang": {
                "zh": {
                    "raw_body": "这是中文原稿第一章。\n\n技术架构坚实稳固。"
                },
                "en": {
                    "raw_body": "This is the English translation chapter 1.\n\nThe technical architecture is robust."
                }
            }
        }
    ]
    book_meta = {
        "title": "双语出版实践",
        "author": "Illacme 架构组",
        "publisher": "Illacme Press",
        "polyglot_langs": ["zh", "en"]
    }

    with tempfile.TemporaryDirectory() as tmpdir:
        out_file = os.path.join(tmpdir, "bilingual_test.docx")
        success = adapter.bind_book(
            manuscript_tree=manuscript_tree,
            book_metadata=book_meta,
            target_lang="zh",
            output_file_path=out_file
        )
        assert success is True
        assert os.path.exists(out_file)
        assert os.path.getsize(out_file) > 1000

        # 打开生成的 docx 验证包含封面、目录、章节对照与版权表格
        exported_doc = docx.Document(out_file)
        assert len(exported_doc.tables) >= 4  # 封面表 + 目录表 + 正文章节表 + 版权声明表
        for tbl in exported_doc.tables:
            assert len(tbl.columns) == 2


def test_docx_polyglot_html_tags_sanitization():
    """验证包含复杂 HTML 标签 (如 home-hero-container 等) 的文本在进入 Word 表格时被 100% 清洗脱壳"""
    doc = docx.Document()
    dirty_html_zh = (
        '<div class="home-hero-container" style="text-align: center; padding: 4rem;">'
        '<div class="hero-badge" style="display: flex;"><span>🧬</span> V50.3 主权全球出版发行中枢已就绪</div>'
        '<h1 class="hero-main-title" style="font-size: 3.8rem;">您的主权化全球出版发行中枢</h1>'
        '</div>'
    )
    dirty_html_en = (
        '<div class="home-hero-container" style="text-align: center; padding: 4rem;">'
        '<div class="hero-badge" style="display: flex;"><span>🧬</span> V50.3 Sovereign Global Publishing Hub Ready</div>'
        '<h1 class="hero-main-title" style="font-size: 3.8rem;">Your Sovereign Global Publishing Hub</h1>'
        '</div>'
    )
    ch = {
        "title": "测试清洗",
        "chapter_by_lang": {
            "zh": {"raw_body": dirty_html_zh},
            "en": {"raw_body": dirty_html_en}
        }
    }
    DocxPolyglotBuilder.render_polyglot_chapter(doc, ch, ["zh", "en"])
    tbl = doc.tables[0]
    all_cell_texts = [cell.text for row in tbl.rows for cell in row.cells]

    for text in all_cell_texts:
        assert "<div" not in text
        assert "style=" not in text
        assert "<h1" not in text
        assert "<span" not in text

    # 验证真实人类文本被完整保留
    combined = " ".join(all_cell_texts)
    assert "V50.3 主权全球出版发行中枢已就绪" in combined
    assert "Your Sovereign Global Publishing Hub" in combined


def test_docx_polyglot_scaffold_cover_toc_colophon():
    """验证多语言骨架生成器 (DocxPolyglotScaffold) 的封面、目录与版权分栏表格"""
    from core.bindery.docx_polyglot_scaffold import DocxPolyglotScaffold
    doc = docx.Document()
    meta = {
        "title": "数字出版集",
        "titles_by_lang": {"zh": "数字出版集", "en": "Digital Publishing Collection"},
        "description": "多语出版实践"
    }
    langs = ["zh", "en"]
    DocxPolyglotScaffold.render_polyglot_cover(doc, meta, langs, "数字出版集", "作者", "出版社")
    assert len(doc.tables) == 1
    assert len(doc.tables[0].columns) == 2
    assert "Digital Publishing Collection" in doc.tables[0].cell(0, 1).text

    tree = [{"id": "ch_1", "title": "第 1 章", "titles_by_lang": {"zh": "第 1 章 绪论", "en": "Chapter 1 Intro"}}]
    DocxPolyglotScaffold.render_polyglot_toc(doc, tree, langs)
    assert len(doc.tables) == 2
    assert "Chapter 1 Intro" in doc.tables[1].cell(1, 1).text

    DocxPolyglotScaffold.render_polyglot_colophon(doc, tree, meta, langs)
    assert len(doc.tables) == 3
    assert len(doc.tables[2].columns) == 2
