# -*- coding: utf-8 -*-
"""
Illacme Plenipes - Test Suite for DOCX Egress & A4 Web Reviewer
测试目标：
1. 验证 DocxBookAdapter 导出的目录大纲消除双重序号，并保持小节缩进
2. 验证 Markdown 链接成功编译为 OpenXML 原生可交互超链接
3. 验证 Markdown 表格编译为原生 Word 表格，水平线编译为分隔线
4. 验证 render_docx_reader_html 仿真软分页与表格/分割线 HTML 还原
🛡️ [SOP-01 规范]：单文件严格 <= 300 行。
"""

import os
import pytest
from core.adapters.egress.ebook.docx import DocxBookAdapter
from core.adapters.egress.ebook.docx_reader import render_docx_reader_html


@pytest.fixture
def sample_manuscript():
    return [
        {
            "id": "ch_1",
            "title": "数字出版主权实践",
            "raw_body": (
                "这是正文第一章的内容，核心路径为 `vault_root` 参数。\n\n"
                "欢迎访问 [⚡ 5 分钟极速上手](./docs/quick-start.html) 获取指引。\n\n"
                "```python\ndef publish():\n    return 'Sovereign Press'\n```\n\n"
                "---\n\n"
                "> 出版即主权。\n\n"
                "- 特性 A\n- 特性 B\n"
            ),
            "headings": [{"title": "为什么选择主权架构", "level": 2}],
        },
        {
            "id": "ch_2",
            "title": "装订中枢工程架构",
            "raw_body": (
                "这是正文第二章的内容。\n\n"
                "详细参考 [📚 查阅官方文档](https://example.com/docs)。\n\n"
                "```mermaid\ngraph TD\n    A[原稿文库] --> B[装订流水线]\n    B --> C[Word 审校印本]\n```\n\n"
                "| 格式 | 引擎 | 职责 |\n"
                "| :--- | :--- | :--- |\n"
                "| Word | docx | 审校印本 |\n"
                "| EPUB | zip  | 移动阅读 |\n\n"
                "---"
            ),
            "headings": [],
        },
    ]


@pytest.fixture
def sample_meta():
    return {
        "title": "Illacme Plenipes 装订典籍",
        "author": "Illacme 架构组",
        "publisher": "Global Press",
        "description": "企业级离线与出版装订解决方案。"
    }


def test_docx_hyperlinks_and_pagination(tmp_path, sample_manuscript, sample_meta):
    out_docx = str(tmp_path / "test_book.docx")
    adapter = DocxBookAdapter()
    from unittest.mock import patch
    fake_png = b"\x89PNG\r\n\x1a\n\x00\x00\x00\rIHDR\x00\x00\x00\x01\x00\x00\x00\x01\x08\x06\x00\x00\x00\x1f\x15c4\x00\x00\x00\nIDATx\x9cc\x00\x01\x00\x00\x05\x00\x01\r\n-\xb4\x00\x00\x00\x00IEND\xaeB`\x82"
    with patch("core.bindery.mermaid_renderer.MermaidRenderer.render_to_png", return_value=fake_png):
        success = adapter.bind_book(
            manuscript_tree=sample_manuscript,
            book_metadata=sample_meta,
            target_lang="zh",
            output_file_path=out_docx
        )
    assert success is True
    assert os.path.exists(out_docx)

    # 1. 验证底层 docx 目录中没有使用 List Number 样式（防止产生 1. 1. 双重序号）
    import docx
    doc = docx.Document(out_docx)
    toc_p = [p for p in doc.paragraphs if "1. 数字出版主权实践" in p.text]
    assert len(toc_p) == 1
    assert toc_p[0].style.name != "List Number"
    assert toc_p[0].text.strip() == "1. 数字出版主权实践"

    # 2. 验证 Word 底层生成了原生 OpenXML 超链接
    hyperlinks = []
    for p in doc.paragraphs:
        for hl in p._p.xpath(".//w:hyperlink"):
            rid = hl.get("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id")
            target = p.part.rels[rid].target_ref if (rid and rid in p.part.rels) else ""
            txt = hl.xpath("string(.)").strip()
            hyperlinks.append((txt, target))

    assert len(hyperlinks) >= 2
    assert any("5 分钟极速上手" in h[0] and "./docs/quick-start.html" in h[1] for h in hyperlinks)
    assert any("查阅官方文档" in h[0] and "https://example.com/docs" in h[1] for h in hyperlinks)

    # 3. 验证 Word 底层生成了原生 Table、代码卡片、水平分割线与 Mermaid 原生图表
    assert len(doc.tables) == 2  # 1 python code card + 1 markdown data table
    code_cards = [t for t in doc.tables if len(t.rows) == 1 and len(t.columns) == 1]
    assert len(code_cards) == 1
    assert "F8FAFC" in code_cards[0].cell(0, 0)._tc.xml  # 浅灰底纹
    assert "10B981" in code_cards[0].cell(0, 0)._tc.xml  # 绿色边框

    # 验证 Mermaid 在 Word 中被成功直接渲染并嵌入为原生图片
    drawings = [p for p in doc.paragraphs if p._p.xpath(".//w:drawing")]
    assert len(drawings) >= 1

    data_tables = [t for t in doc.tables if len(t.rows) > 1]
    assert len(data_tables) == 1
    table_headers = [c.text.strip() for c in data_tables[0].rows[0].cells]
    assert table_headers == ["格式", "引擎", "职责"]
    hr_paragraphs = [p for p in doc.paragraphs if p._p.xpath(".//w:pBdr")]
    assert len(hr_paragraphs) >= 1

    # 4. 验证 Web 审校渲染器还原了可交互超链接、Mermaid 图码双模切换与表格
    html_out = render_docx_reader_html(out_docx)
    assert '<article class="doc-page doc-cover-page" id="page-1">' in html_out
    assert "Illacme Plenipes 装订典籍" in html_out
    assert 'class="doc-link"' in html_out
    assert 'href="./docs/quick-start.html"' in html_out
    assert 'href="https://example.com/docs"' in html_out
    assert '<div class="doc-code-block">' in html_out
    assert '<span class="doc-code-lang">PYTHON</span>' in html_out
    assert '<code class="doc-inline-code">vault_root</code>' in html_out
    assert '<div class="doc-table-wrap"><table>' in html_out
    assert '<th>格式</th>' in html_out
    assert '<hr class="doc-hr">' in html_out

    # 5. 验证离线脚本支持与 Word 原生嵌入图表 (如 Mermaid) 成功还原为图片
    assert '<script src="/dashboard/vendor/mermaid.min.js"></script>' in html_out
    assert '<img src="data:image/' in html_out
    assert 'class="doc-embedded-img"' in html_out

    # 6. 验证多页容器与分页存在
    assert '<div class="doc-pages-container" id="pages-container">' in html_out
    assert 'id="page-2"' in html_out
    assert "📖 目录大纲 (Contents)" in html_out

    # 7. 验证快速跳转首页、尾页、指定页与移动端响应式支持
    assert 'goToPage(1)' in html_out
    assert 'goToPage(' in html_out and '尾页' in html_out
    assert 'id="page-jump-input"' in html_out
    assert 'class="viewer-pagination"' in html_out
    assert '@media (max-width: 768px)' in html_out
    assert '@media (max-width: 480px)' in html_out


def test_docx_mermaid_fallback_to_code_card(tmp_path, sample_manuscript, sample_meta):
    """验证无无头渲染引擎环境时，Mermaid 平滑降级为 Word 代码卡片"""
    from unittest.mock import patch
    out_docx = str(tmp_path / "fallback_book.docx")
    adapter = DocxBookAdapter()
    with patch("core.bindery.mermaid_renderer.MermaidRenderer.render_to_png", return_value=None):
        success = adapter.bind_book(
            manuscript_tree=sample_manuscript,
            book_metadata=sample_meta,
            target_lang="zh",
            output_file_path=out_docx
        )
    assert success is True
    import docx
    doc = docx.Document(out_docx)
    # 降级模式下：1 个 python 代码卡片 + 1 个 mermaid 降级代码卡片 + 1 个 markdown 数据表格 = 3
    assert len(doc.tables) == 3
    code_cards = [t for t in doc.tables if len(t.rows) == 1 and len(t.columns) == 1]
    assert len(code_cards) == 2
    # 无原生 drawing 图片
    drawings = [p for p in doc.paragraphs if p._p.xpath(".//w:drawing")]
    assert len(drawings) == 0

