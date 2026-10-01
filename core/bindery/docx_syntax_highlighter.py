# -*- coding: utf-8 -*-
"""
Illacme Plenipes - DOCX Code Block Syntax Highlighter
模块职责：基于 Pygments 词法分析器，将多语言代码块转译为 Word 原生带高亮着色与样式的 Run 节点。
适用场景：Word (.docx) 审校本导出，提供出版级 Monokai Light / GitHub Light 风格的语法着色。
🛡️ [SOP-01 规范]：单文件严格 <= 300 行。
"""

from typing import Any, Tuple, List, Optional
from docx.shared import Pt, RGBColor

try:
    import pygments
    from pygments.lexers import get_lexer_by_name, guess_lexer, TextLexer
    from pygments.token import Token
    _HAS_PYGMENTS = True
except ImportError:
    _HAS_PYGMENTS = False


class DocxSyntaxHighlighter:
    """Word 原生代码块语法高亮着色器"""

    # 浅色出版物高对比度语法配色表 (RGB)
    COLOR_KEYWORD = RGBColor(124, 58, 237)      # #7C3AED 紫罗兰蓝 (关键字)
    COLOR_NAME_FN = RGBColor(2, 132, 199)       # #0284C7 沉稳天蓝 (函数名/类名)
    COLOR_BUILTIN = RGBColor(130, 80, 223)      # #8250DF 靛蓝 (内置函数/装饰器)
    COLOR_STRING = RGBColor(21, 128, 61)        # #15803D 森林绿 (字符串)
    COLOR_COMMENT = RGBColor(100, 116, 139)     # #64748B 沉稳灰 (注释)
    COLOR_NUMBER = RGBColor(217, 119, 6)        # #D97706 琥珀橙 (数字)
    COLOR_OPERATOR = RGBColor(51, 65, 85)       # #334155 深碳灰 (运算符/标点)
    COLOR_DEFAULT = RGBColor(30, 41, 59)        # #1E293B 炭黑 (普通标识符与正文)

    @classmethod
    def _resolve_token_style(cls, ttype: Any) -> Tuple[RGBColor, bool, bool]:
        """根据 Pygments Token 类型推导 (RGB颜色, 是否加粗, 是否斜体)"""
        if ttype in Token.Keyword:
            return cls.COLOR_KEYWORD, True, False
        if ttype in (Token.Name.Function, Token.Name.Class):
            return cls.COLOR_NAME_FN, True, False
        if ttype in (Token.Name.Builtin, Token.Name.Decorator, Token.Name.Exception):
            return cls.COLOR_BUILTIN, False, False
        if ttype in Token.String:
            return cls.COLOR_STRING, False, False
        if ttype in Token.Comment:
            return cls.COLOR_COMMENT, False, True
        if ttype in Token.Number:
            return cls.COLOR_NUMBER, False, False
        if ttype in (Token.Operator, Token.Punctuation):
            return cls.COLOR_OPERATOR, False, False
        return cls.COLOR_DEFAULT, False, False

    @classmethod
    def _get_lexer(cls, code_text: str, lang: str = ""):
        """获取匹配的 Pygments 词法分析器"""
        if not _HAS_PYGMENTS:
            return None
        lang_clean = (lang or "").strip().lower()
        if lang_clean and lang_clean not in ("text", "txt", "plain"):
            try:
                return get_lexer_by_name(lang_clean, stripall=False)
            except Exception:
                pass
        try:
            return guess_lexer(code_text)
        except Exception:
            return TextLexer()

    @classmethod
    def highlight_code_block(
        cls,
        paragraph: Any,
        code_text: str,
        lang: str = "",
        font_name: str = "Consolas",
        font_size_pt: float = 9.5
    ) -> None:
        """
        向 Word 段落流式添加具备出版级语法高亮着色的 Run 节点。
        
        Args:
            paragraph: docx 段落对象
            code_text: 代码块纯文本内容
            lang: 编程语言标识 (如 python, javascript, rust, bash 等)
            font_name: 等宽字体名称
            font_size_pt: 字体大小 (pt)
        """
        if not code_text:
            return

        lexer = cls._get_lexer(code_text, lang)
        if not lexer:
            # 降级模式：直接注入单一等宽纯文本 Run
            r = paragraph.add_run(code_text)
            r.font.name = font_name
            r.font.size = Pt(font_size_pt)
            try:
                r.font.color.rgb = cls.COLOR_DEFAULT
            except Exception:
                pass
            return

        try:
            tokens = list(pygments.lex(code_text, lexer))
            # 性能优化：将具有相同样式的连续 token 合并，缩减 Word XML 节点数
            merged_spans: List[Tuple[str, RGBColor, bool, bool]] = []
            curr_text, curr_color, curr_bold, curr_italic = "", None, False, False

            for ttype, val in tokens:
                if not val:
                    continue
                color, bold, italic = cls._resolve_token_style(ttype)
                if curr_color == color and curr_bold == bold and curr_italic == italic:
                    curr_text += val
                else:
                    if curr_text:
                        merged_spans.append((curr_text, curr_color, curr_bold, curr_italic))
                    curr_text, curr_color, curr_bold, curr_italic = val, color, bold, italic

            if curr_text:
                merged_spans.append((curr_text, curr_color, curr_bold, curr_italic))

            # 将合并后的样式块逐一添加至段落
            for text_chunk, color, bold, italic in merged_spans:
                r = paragraph.add_run(text_chunk)
                r.font.name = font_name
                r.font.size = Pt(font_size_pt)
                if bold:
                    r.bold = True
                if italic:
                    r.italic = True
                try:
                    r.font.color.rgb = color
                except Exception:
                    pass

        except Exception:
            # 防御性自愈降级
            r = paragraph.add_run(code_text)
            r.font.name = font_name
            r.font.size = Pt(font_size_pt)
            try:
                r.font.color.rgb = cls.COLOR_DEFAULT
            except Exception:
                pass
