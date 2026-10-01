# -*- coding: utf-8 -*-
"""
Illacme Plenipes - DOCX Polyglot Scaffold (Cover, TOC & Colophon)
模块职责：为 Word (.docx) 多语言对照版构建并排分栏封面、并排分栏目录大纲与并排版权声明页。
适用场景：多语言对照 Word 印本导出，确保封面、目录与版权页遵循全书双栏/多栏对照视觉体系。
🛡️ [SOP-01 规范]：单文件严格 <= 300 行。
"""

import os
from datetime import datetime
from typing import Any, Dict, List, Optional
from docx.oxml import parse_xml
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from core.adapters.egress.ebook.colophon import ColophonBuilder
from core.bindery.docx_text_parser import DocxTextParser


class DocxPolyglotScaffold:
    """🌐 Word 多语言对照版骨架装帧组件（封面、目录、版权页）"""

    LANG_META = {
        "zh": ("🇨🇳 简体中文", "ZH"),
        "en": ("🇬🇧 English Translation", "EN"),
        "ja": ("🇯🇵 日本語訳", "JA"),
        "fr": ("🇫🇷 Traduction Française", "FR"),
        "de": ("🇩🇪 Deutsche Übersetzung", "DE"),
        "es": ("🇪🇸 Traducción Española", "ES"),
        "ru": ("🇷🇺 Русский перевод", "RU")
    }

    @classmethod
    def render_polyglot_cover(
        cls,
        doc: Any,
        book_metadata: Dict[str, Any],
        polyglot_langs: List[str],
        default_title: str,
        default_author: str,
        default_publisher: str
    ) -> None:
        """为 Word 文档构建多语言分栏并排封面"""
        cols_count = len(polyglot_langs)
        titles_by_lang = book_metadata.get("titles_by_lang", {})

        # 1. 顶部出品品牌徽章
        p_top = doc.add_paragraph()
        p_top.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_top.paragraph_format.space_before = Pt(40)
        p_top.paragraph_format.space_after = Pt(20)
        r_pub = p_top.add_run(default_publisher.upper())
        r_pub.font.size = Pt(10)
        r_pub.font.bold = True
        try: r_pub.font.color.rgb = RGBColor(16, 185, 129)
        except Exception: pass

        # 2. 中部多语并排分栏书名表格
        tbl = doc.add_table(rows=1, cols=cols_count)
        tbl.style = "Table Grid"

        for c_idx, l in enumerate(polyglot_langs):
            cell = tbl.cell(0, c_idx)
            cell.text = ""
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(20)
            p.paragraph_format.space_after = Pt(20)

            lang_name, tag = cls.LANG_META.get(l, (l.upper(), l.upper()))
            r_badge = p.add_run(f"[{tag}] {lang_name}\n\n")
            r_badge.font.size = Pt(9.5)
            r_badge.bold = True
            try: r_badge.font.color.rgb = RGBColor(2, 132, 199)
            except Exception: pass

            l_title = titles_by_lang.get(l) or default_title
            r_title = p.add_run(f"{l_title}\n\n")
            r_title.font.size = Pt(20)
            r_title.bold = True
            try: r_title.font.color.rgb = RGBColor(15, 23, 42)
            except Exception: pass

            desc = book_metadata.get("description", "")
            if desc:
                r_desc = p.add_run(desc)
                r_desc.font.size = Pt(9.5)
                r_desc.italic = True
                try: r_desc.font.color.rgb = RGBColor(100, 116, 139)
                except Exception: pass

            # 边框设置：全透明无边框，左右内边距
            try:
                tcPr = cell._tc.get_or_add_tcPr()
                tcPr.append(parse_xml(
                    '<w:tcBorders xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
                    '<w:top w:val="none"/><w:left w:val="none"/><w:right w:val="none"/><w:bottom w:val="none"/>'
                    '</w:tcBorders>'
                ))
                tcPr.append(parse_xml(
                    '<w:tcMar xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
                    '<w:left w:w="160" w:type="dxa"/><w:right w:w="160" w:type="dxa"/>'
                    '</w:tcMar>'
                ))
            except Exception:
                pass

        # 3. 底部责任信息与规格徽标
        p_bot = doc.add_paragraph()
        p_bot.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_bot.paragraph_format.space_before = Pt(80)
        p_bot.paragraph_format.space_after = Pt(20)
        lang_tags = " · ".join([l.upper() for l in polyglot_langs])
        r_bot = p_bot.add_run(
            f"责任作者：{default_author}\n"
            f"装订日期：{datetime.now().strftime('%Y-%m-%d')} · 双语矩阵规格：{lang_tags} PARALLEL EDITION"
        )
        r_bot.font.size = Pt(9.5)
        try: r_bot.font.color.rgb = RGBColor(71, 85, 105)
        except Exception: pass

        doc.add_page_break()

    @classmethod
    def render_polyglot_toc(
        cls,
        doc: Any,
        manuscript_tree: List[Dict[str, Any]],
        polyglot_langs: List[str]
    ) -> None:
        """为 Word 文档构建多语言分栏并排目录大纲"""
        cols_count = len(polyglot_langs)
        doc.add_heading("📖 目录大纲 · Contents", level=1)

        tbl = doc.add_table(rows=len(manuscript_tree) + 1, cols=cols_count)
        tbl.style = "Table Grid"

        # 表头行：各语言徽章
        hdr_row = tbl.rows[0]
        try:
            hdr_row._tr.get_or_add_trPr().append(parse_xml('<w:tblHeader xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"/>'))
        except Exception:
            pass

        for c_idx, l in enumerate(polyglot_langs):
            cell = hdr_row.cells[c_idx]
            cell.text = ""
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(4)
            lang_name, tag = cls.LANG_META.get(l, (l.upper(), l.upper()))
            r = p.add_run(f"[{tag}] {lang_name} 目录")
            r.bold = True
            r.font.size = Pt(9.5)
            try: r.font.color.rgb = RGBColor(2, 132, 199)
            except Exception: pass

            try:
                tcPr = cell._tc.get_or_add_tcPr()
                tcPr.append(parse_xml('<w:shd xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" w:fill="F1F5F9"/>'))
                tcPr.append(parse_xml(
                    '<w:tcBorders xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
                    '<w:top w:val="none"/><w:left w:val="none"/><w:right w:val="none"/>'
                    '<w:bottom w:val="single" w:sz="12" w:space="0" w:color="0284C7"/>'
                    '</w:tcBorders>'
                ))
            except Exception: pass

        # 表体行：各章节超链接水平对齐
        for r_idx, ch in enumerate(manuscript_tree):
            row = tbl.rows[r_idx + 1]
            ch_id = ch.get("id") or f"ch_{r_idx + 1}"
            titles_by_lang = ch.get("titles_by_lang", {})

            for c_idx, l in enumerate(polyglot_langs):
                cell = row.cells[c_idx]
                cell.text = ""
                p = cell.paragraphs[0]
                p.paragraph_format.space_before = Pt(3)
                p.paragraph_format.space_after = Pt(3)

                ch_title = titles_by_lang.get(l) or ch.get("title", f"第 {r_idx + 1} 章")
                DocxTextParser.add_hyperlink(p, f"#{ch_id}", f"{r_idx + 1}. {ch_title}", color="1e293b", underline=False, bold=True, font_size_pt=10.0)

                # 单元格边框：无竖线，淡灰底线
                try:
                    tcPr = cell._tc.get_or_add_tcPr()
                    tcPr.append(parse_xml(
                        '<w:tcBorders xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
                        '<w:top w:val="none"/><w:left w:val="none"/><w:right w:val="none"/>'
                        '<w:bottom w:val="single" w:sz="4" w:space="0" w:color="F1F5F9"/>'
                        '</w:tcBorders>'
                    ))
                    tcPr.append(parse_xml(
                        '<w:tcMar xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
                        '<w:left w:w="100" w:type="dxa"/><w:right w:w="100" w:type="dxa"/>'
                        '</w:tcMar>'
                    ))
                except Exception: pass

        doc.add_page_break()

    @classmethod
    def render_polyglot_colophon(
        cls,
        doc: Any,
        manuscript_tree: List[Dict[str, Any]],
        book_metadata: Dict[str, Any],
        polyglot_langs: List[str]
    ) -> None:
        """为 Word 文档构建多语言分栏并排版权声明页"""
        doc.add_page_break()
        doc.add_heading("📜 版权声明 · Colophon", level=1)
        cols_count = len(polyglot_langs)

        tbl = doc.add_table(rows=2, cols=cols_count)
        tbl.style = "Table Grid"

        # 表头
        hdr_row = tbl.rows[0]
        for c_idx, l in enumerate(polyglot_langs):
            cell = hdr_row.cells[c_idx]
            cell.text = ""
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(4)
            lang_name, tag = cls.LANG_META.get(l, (l.upper(), l.upper()))
            r = p.add_run(f"[{tag}] {lang_name} 出版版权")
            r.bold = True
            r.font.size = Pt(9.5)
            try: r.font.color.rgb = RGBColor(2, 132, 199)
            except Exception: pass
            try:
                tcPr = cell._tc.get_or_add_tcPr()
                tcPr.append(parse_xml('<w:shd xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" w:fill="F1F5F9"/>'))
                tcPr.append(parse_xml(
                    '<w:tcBorders xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
                    '<w:top w:val="none"/><w:left w:val="none"/><w:right w:val="none"/>'
                    '<w:bottom w:val="single" w:sz="12" w:space="0" w:color="0284C7"/>'
                    '</w:tcBorders>'
                ))
            except Exception: pass

        # 表体各语种版权数据卡片
        body_row = tbl.rows[1]
        for c_idx, l in enumerate(polyglot_langs):
            cell = body_row.cells[c_idx]
            cell.text = ""
            c_dict = ColophonBuilder.build_colophon_data(manuscript_tree=manuscript_tree, book_metadata=book_metadata, format_name="docx")
            t_by_l = book_metadata.get("titles_by_lang", {})
            cur_t = t_by_l.get(l) or c_dict.get("title", "数字典籍")

            is_zh = (l == "zh")
            lines = [
                f"{'出版典籍' if is_zh else 'Title'}: {cur_t}",
                f"{'责任作者' if is_zh else 'Author'}: {c_dict.get('author', 'Illacme Team')}",
                f"{'出版机构' if is_zh else 'Publisher'}: {c_dict.get('publisher', 'Illacme Press')}",
                f"{'装订日期' if is_zh else 'Date'}: {datetime.now().strftime('%Y-%m-%d')}",
                f"{'技术驱动' if is_zh else 'Engine'}: Illacme Plenipes Sovereign Digital Bindery Hub",
                f"{'版权规范' if is_zh else 'License'}: {c_dict.get('license', 'All Rights Reserved')}"
            ]

            for line_idx, line in enumerate(lines):
                p = cell.paragraphs[0] if line_idx == 0 else cell.add_paragraph()
                p.paragraph_format.space_before = Pt(2)
                p.paragraph_format.space_after = Pt(2)
                r = p.add_run(line)
                r.font.size = Pt(9.0)
                try: r.font.color.rgb = RGBColor(71, 85, 105)
                except Exception: pass

            try:
                tcPr = cell._tc.get_or_add_tcPr()
                tcPr.append(parse_xml(
                    '<w:tcBorders xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
                    '<w:top w:val="none"/><w:left w:val="none"/><w:right w:val="none"/><w:bottom w:val="none"/>'
                    '</w:tcBorders>'
                ))
                tcPr.append(parse_xml(
                    '<w:tcMar xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
                    '<w:top w:w="120" w:type="dxa"/><w:left w:w="120" w:type="dxa"/>'
                    '<w:bottom w:w="120" w:type="dxa"/><w:right w:w="120" w:type="dxa"/>'
                    '</w:tcMar>'
                ))
            except Exception: pass
