# -*- coding: utf-8 -*-
"""
Illacme Plenipes - DOCX Rich Text & Hyperlink Tokenizer
模块职责：将 Markdown 行内超链接、数学公式、加粗、原生表格与水平线解析为 Word 节点。
适用场景：出版级 Word 审校本导出，确保超链接、表格与分割线在 Word 与 WPS 中交互可点。
🛡️ [SOP-01 规范]：单文件严格 <= 300 行。
"""

import re
import html
from typing import Any, List
import docx.opc.constants
from docx.oxml import parse_xml


class DocxTextParser:
    """Word 文档富文本、超链接与原生表格/分割线构建器"""

    @staticmethod
    def add_hyperlink(paragraph: Any, url: str, text: str, color: str = "2563eb", underline: bool = True, bold: bool = False, font_size_pt: float = None) -> None:
        """向 Word 段落添加原生交互式超链接节点（支持外部 URL 与内部书签锚点）"""
        if not text or not url:
            if text:
                r = paragraph.add_run(text)
                if bold: r.bold = True
                if font_size_pt:
                    from docx.shared import Pt
                    r.font.size = Pt(font_size_pt)
            return

        clean_u = url.strip()
        m_anc = re.search(r"#(ch_\d+|[A-Za-z0-9_-]+)$", clean_u)
        m_ch = re.match(r"^(?:text/)?(ch_\d+)(?:\.xhtml)?(?:#(.*?))?$", clean_u)
        if clean_u.startswith("#"):
            anchor = clean_u[1:].strip()
        elif m_anc:
            anchor = m_anc.group(1)
        elif m_ch:
            anchor = m_ch.group(2) or m_ch.group(1)
        else:
            anchor = None

        u_attr = ' w:val="single"' if underline else ""
        b_tag = '<w:b/>' if bold else ""
        sz_tag = f'<w:sz w:val="{int(round(font_size_pt * 2))}"/><w:szCs w:val="{int(round(font_size_pt * 2))}"/>' if font_size_pt else ""
        esc_t = html.escape(text)

        try:
            if anchor:
                xml_str = (
                    f'<w:hyperlink xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" w:anchor="{anchor}" w:history="1">'
                    f'<w:r><w:rPr><w:color w:val="{color}"/>{b_tag}{sz_tag}<w:u{u_attr}/></w:rPr><w:t>{esc_t}</w:t></w:r>'
                    f'</w:hyperlink>'
                )
            else:
                part = paragraph.part
                r_id = part.relate_to(clean_u, docx.opc.constants.RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
                xml_str = (
                    f'<w:hyperlink xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
                    f'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" r:id="{r_id}">'
                    f'<w:r><w:rPr><w:color w:val="{color}"/>{b_tag}{sz_tag}<w:u{u_attr}/></w:rPr><w:t>{esc_t}</w:t></w:r>'
                    f'</w:hyperlink>'
                )
            paragraph._p.append(parse_xml(xml_str))
        except Exception:
            r = paragraph.add_run(text)
            r.font.underline = underline
            if bold: r.bold = True
            if font_size_pt:
                from docx.shared import Pt
                r.font.size = Pt(font_size_pt)

    @classmethod
    def append_rich_runs(cls, paragraph: Any, text: str, is_italic: bool = False, is_bold: bool = False) -> None:
        """将包含超链接、公式与加粗的 Markdown 文本分词并流式注入 Word 段落"""
        if not text:
            return
        split_pattern = re.compile(r"(\[[^\]]+\]\([^)]+\)|(?<![\w\\\$])\$(?!\s|\d)[^\$\n]+?(?<!\s)\$(?![\w\$]))")
        parts = split_pattern.split(text)

        for part in parts:
            if not part:
                continue
            m_link = re.match(r"^\[([^\]]+)\]\(([^)]+)\)$", part)
            if m_link:
                cls.add_hyperlink(paragraph, m_link.group(2).strip(), m_link.group(1).strip(), color="2563eb", underline=True)
                continue

            if part.startswith("$") and part.endswith("$") and len(part) > 2 and not re.match(r"^\$\s*[\d,]+(?:\.\d+)?\s*\$$", part):
                latex = part[1:-1].strip()
                from core.bindery.math_packager import MathPackager
                omml = MathPackager.latex_to_omml(latex, display="inline")
                if omml:
                    try:
                        paragraph._p.append(parse_xml(omml))
                        continue
                    except Exception:
                        pass
                r = paragraph.add_run(part)
                if is_italic:
                    r.font.italic = True
                if is_bold:
                    r.font.bold = True
                continue

            cls._append_bold_and_code_segments(paragraph, part, is_italic=is_italic, is_bold=is_bold)

    @classmethod
    def _append_bold_and_code_segments(cls, paragraph: Any, text: str, is_italic: bool = False, is_bold: bool = False) -> None:
        """解析文本中的加粗 **...** 与行内代码 `...` 并注入等宽浅灰胶囊底纹"""
        from docx.shared import Pt, RGBColor
        sub_parts = re.split(r"(\*\*[^*]+?\*\*|`[^`]+?`)", text)
        for sp in sub_parts:
            if not sp: continue
            if sp.startswith("**") and sp.endswith("**") and len(sp) > 4:
                r = paragraph.add_run(sp[2:-2])
                r.bold = True
                if is_italic: r.font.italic = True
            elif sp.startswith("`") and sp.endswith("`") and len(sp) > 2:
                r = paragraph.add_run(sp[1:-1])
                r.font.name = "Consolas"
                r.font.size = Pt(9.5)
                try:
                    r.font.color.rgb = RGBColor(15, 23, 42)
                    r._r.get_or_add_rPr().append(parse_xml('<w:shd xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" w:fill="F1F5F9"/>'))
                except Exception: pass
            else:
                clean_t = re.sub(r"[*_]", "", sp)
                if clean_t:
                    r = paragraph.add_run(clean_t)
                    if is_bold: r.font.bold = True
                    if is_italic: r.font.italic = True

    @classmethod
    def create_code_block_card(cls, doc: Any, code_lines: List[str], lang: str = "") -> None:
        """在 Word 中创建出版级代码块容器卡片（1x1 浅灰底纹单格表格，带强调色左边框）"""
        if not code_lines and not lang: return
        from docx.shared import Pt, RGBColor
        tbl = doc.add_table(rows=1, cols=1)
        tbl.style = "Table Grid"
        cell = tbl.cell(0, 0)
        cell.text = ""

        try:
            tcPr = cell._tc.get_or_add_tcPr()
            tcPr.append(parse_xml('<w:shd xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" w:fill="F8FAFC"/>'))
            tcPr.append(parse_xml(
                '<w:tcBorders xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
                '<w:top w:val="single" w:sz="4" w:space="0" w:color="E2E8F0"/>'
                '<w:left w:val="single" w:sz="18" w:space="0" w:color="10B981"/>'
                '<w:bottom w:val="single" w:sz="4" w:space="0" w:color="E2E8F0"/>'
                '<w:right w:val="single" w:sz="4" w:space="0" w:color="E2E8F0"/>'
                '</w:tcBorders>'
            ))
            tcPr.append(parse_xml(
                '<w:tcMar xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
                '<w:top w:w="120" w:type="dxa"/><w:left w:w="180" w:type="dxa"/>'
                '<w:bottom w:w="120" w:type="dxa"/><w:right w:w="180" w:type="dxa"/>'
                '</w:tcMar>'
            ))
        except Exception: pass

        p_cur = cell.paragraphs[0]
        if lang and lang.lower() not in ("text", "txt", "plain"):
            p_cur.paragraph_format.space_before = 0
            p_cur.paragraph_format.space_after = 25400  # 2pt
            r_lang = p_cur.add_run(lang.upper())
            r_lang.font.name = "Consolas"
            r_lang.font.size = Pt(8.5)
            r_lang.bold = True
            try: r_lang.font.color.rgb = RGBColor(100, 116, 139)
            except Exception: pass
            p_cur = cell.add_paragraph()

        p_cur.paragraph_format.space_before = 0
        p_cur.paragraph_format.space_after = 0
        p_cur.paragraph_format.line_spacing = 1.15
        from core.bindery.docx_syntax_highlighter import DocxSyntaxHighlighter
        DocxSyntaxHighlighter.highlight_code_block(p_cur, "\n".join(code_lines), lang=lang, font_name="Consolas", font_size_pt=9.5)

        p_gap = doc.add_paragraph()
        p_gap.paragraph_format.space_before = 0
        p_gap.paragraph_format.space_after = 38100

    @classmethod
    def parse_markdown_table(cls, doc: Any, table_lines: List[str]) -> None:
        """将 Markdown 表格文本列表编译为 Word 原生带样式网格表格"""
        if not table_lines:
            return
        rows_data = []
        for line in table_lines:
            content = line.strip().strip("|")
            if re.match(r"^[\s\-:|]+$", content):
                continue
            rows_data.append([c.strip() for c in content.split("|")])

        if not rows_data:
            return
        cols_count = max(len(r) for r in rows_data)
        tbl = doc.add_table(rows=len(rows_data), cols=cols_count)
        tbl.style = "Table Grid"

        for r_idx, r_cells in enumerate(rows_data):
            for c_idx in range(cols_count):
                val = r_cells[c_idx] if c_idx < len(r_cells) else ""
                cell = tbl.cell(r_idx, c_idx)
                cell.text = ""
                p = cell.paragraphs[0]
                p.paragraph_format.space_before = 25400  # 2pt
                p.paragraph_format.space_after = 25400   # 2pt
                if r_idx == 0:
                    try:
                        shd = parse_xml('<w:shd xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" w:fill="F1F5F9"/>')
                        cell._tc.get_or_add_tcPr().append(shd)
                    except Exception:
                        pass
                    cls.append_rich_runs(p, val, is_bold=True)
                else:
                    cls.append_rich_runs(p, val)

    @staticmethod
    def add_horizontal_rule(doc: Any) -> None:
        """向 Word 文档插入标准细灰水平分割线"""
        p = doc.add_paragraph()
        p.paragraph_format.space_before = 76200
        p.paragraph_format.space_after = 76200
        try:
            pPr = p._p.get_or_add_pPr()
            pBdr = parse_xml('<w:pBdr xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:bottom w:val="single" w:sz="6" w:space="1" w:color="CBD5E1"/></w:pBdr>')
            pPr.append(pBdr)
        except Exception:
            pass

    @classmethod
    def create_mermaid_figure(cls, doc: Any, code_lines: List[str]) -> bool:
        """在 Word 中将 Mermaid 文本编译为居中高清图片嵌入"""
        if not code_lines:
            return False
        try:
            from core.bindery.mermaid_renderer import MermaidRenderer
            png_bytes = MermaidRenderer.render_to_png("\n".join(code_lines))
            if not png_bytes:
                return False
            import io
            from docx.shared import Inches, Pt, RGBColor
            p_img = doc.add_paragraph()
            p_img.alignment = 1  # 居中对齐
            p_img.paragraph_format.space_before = 76200  # 6pt
            p_img.paragraph_format.space_after = 25400   # 2pt
            run_img = p_img.add_run()
            run_img.add_picture(io.BytesIO(png_bytes), width=Inches(5.6))

            p_cap = doc.add_paragraph()
            p_cap.alignment = 1
            p_cap.paragraph_format.space_before = 0
            p_cap.paragraph_format.space_after = 76200
            r_cap = p_cap.add_run("📊 架构与流程示意图 (Mermaid)")
            r_cap.font.name = "Consolas"
            r_cap.font.size = Pt(8.5)
            r_cap.italic = True
            try: r_cap.font.color.rgb = RGBColor(100, 116, 139)
            except Exception: pass
            return True
        except Exception:
            return False

