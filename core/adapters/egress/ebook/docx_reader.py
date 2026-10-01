# -*- coding: utf-8 -*-
"""
Illacme Plenipes - Word Document (.docx) Interactive Web Reviewer
模块职责：将出版级 Word 审校本 (.docx) 物理切分为标准 A4 独立纸张卡片视图。
适用场景：网页端免下载直接在线批阅、多语言分栏封面/目录/版权/对照网格渲染、超链接锚点平滑滚动。
🛡️ [SOP-01 规范]：单文件严格 <= 300 行。
"""

import os, re, html, base64
from typing import Optional, List, Tuple, Any
from core.adapters.egress.ebook.docx_reader_template import assemble_docx_template

PAGE_CAPACITY_LIMIT = 880


def render_docx_reader_html(docx_path: str, filename: Optional[str] = None) -> str:
    """解析 .docx 文件并返回拟真 A4 多页物理装订审阅网页"""
    fn = filename or os.path.basename(docx_path)
    file_size_kb = round(os.path.getsize(docx_path) / 1024, 1) if os.path.exists(docx_path) else 0
    pages_html, total_pages = [], 0

    try:
        import docx
        from docx.text.paragraph import Paragraph
        from docx.table import Table

        doc = docx.Document(docx_path)
        raw_pages: List[List[Tuple[str, Any]]] = []
        cur_items: List[Tuple[str, Any]] = []
        cur_points = 0

        for child in doc._body._element:
            tag = child.tag.split("}")[-1]
            if tag == "p":
                p = Paragraph(child, doc)
                has_pb = bool(p._p.xpath('.//w:br[@w:type="page"]'))
                raw_t = p.text.strip()
                pts = _estimate_p_points(p, raw_t)

                if has_pb:
                    if cur_items: raw_pages.append(cur_items); cur_items = []; cur_points = 0
                    if raw_t: cur_items.append(("p", p)); cur_points += pts
                    continue

                if raw_t or p._p.xpath('.//w:pBdr') or p._p.xpath('.//w:drawing'):
                    if cur_points + pts > PAGE_CAPACITY_LIMIT and cur_items:
                        raw_pages.append(cur_items); cur_items = [("p", p)]; cur_points = pts
                    else:
                        cur_items.append(("p", p)); cur_points += pts
            elif tag == "tbl":
                tbl = Table(child, doc)
                is_code = (len(tbl.rows) == 1 and len(tbl.columns) == 1)
                num_rows = len(tbl.rows)
                if is_code or num_rows <= 3:
                    pts = (20 + max(1, len(tbl.rows[0].cells[0].text.splitlines())) * 17) if is_code else (35 + sum(_estimate_row_pts(r) for r in tbl.rows))
                    if cur_points + pts > PAGE_CAPACITY_LIMIT and cur_items:
                        raw_pages.append(cur_items); cur_items = [("tbl", (tbl, 0, num_rows))]; cur_points = pts
                    else:
                        cur_items.append(("tbl", (tbl, 0, num_rows))); cur_points += pts
                else:
                    # 大表格智能动态切片：按中英段落实际行高自然切页，杜绝强塞压缩
                    r_idx = 0
                    while r_idx < num_rows:
                        rem_pts = PAGE_CAPACITY_LIMIT - cur_points
                        if rem_pts < 180 and cur_items:
                            raw_pages.append(cur_items); cur_items = []; cur_points = 0
                        seg_pts, nxt = 35, r_idx
                        while nxt < num_rows:
                            r_pts = _estimate_row_pts(tbl.rows[nxt])
                            if (cur_points + seg_pts + r_pts > PAGE_CAPACITY_LIMIT) and (nxt > r_idx):
                                break
                            seg_pts += r_pts
                            nxt += 1
                        cur_items.append(("tbl", (tbl, r_idx, nxt)))
                        cur_points += seg_pts
                        r_idx = nxt
                        if cur_points >= PAGE_CAPACITY_LIMIT and r_idx < num_rows:
                            raw_pages.append(cur_items); cur_items = []; cur_points = 0

        if cur_items: raw_pages.append(cur_items)
        total_pages = len(raw_pages)
        for p_idx, page in enumerate(raw_pages, 1):
            pages_html.append(_render_single_page(page, p_idx, total_pages))
    except Exception as e:
        pages_html.append(f'<div class="doc-error">⚠️ 读取 Word 发生异常: {html.escape(str(e))}</div>')
        total_pages = 1

    content_html = "\n".join(pages_html)
    download_url = f"/api/bindery/download?file={html.escape(fn)}"
    return assemble_docx_template(fn, file_size_kb, total_pages, content_html, download_url)


def _estimate_row_pts(row: Any) -> int:
    """估算表格单行占用的版面容量点数（根据双栏/多栏最长文本折算行高）"""
    max_t = max((len(c.text.strip()) for c in row.cells), default=0)
    lines = max(1, max_t // 24 + (1 if max_t % 24 else 0))
    return max(26, min(220, 14 + lines * 18))


def _estimate_p_points(p: Any, text: str) -> int:
    """估算段落所占用的版面容量点数"""
    if not text and p._p.xpath('.//w:pBdr'): return 25
    if not text and p._p.xpath('.//w:drawing'): return 180
    if not text: return 0
    s = (p.style.name if p.style else "").lower()
    if "heading 1" in s: return 60
    if "heading 2" in s: return 42
    if "heading 3" in s: return 32
    if text.startswith(("·", "•")): return 13 + max(0, len(text) // 40) * 14
    if re.match(r"^\d+\.\s+", text): return 18 + max(0, len(text) // 40) * 16
    if text.startswith(("-", "*", ">")): return 16 + max(0, len(text) // 40) * 16
    lines = max(1, len(text) // 36 + (1 if len(text) % 36 else 0))
    return 8 + lines * 22


def _resolve_anchor(u: str) -> Optional[str]:
    m_xh = re.match(r'^(?:text/)?(ch_\d+)\.xhtml(?:#(.*?))?$', u)
    return (m_xh.group(2) or m_xh.group(1)) if m_xh else None


def _render_paragraph_html(p: Any) -> Tuple[str, Optional[str], bool]:
    """将段落解析为 HTML，脱除残余标签并识别超链接与书签锚点"""
    part, chunks, heading_id = p.part, [], None
    for bm in p._p.xpath('.//w:bookmarkStart'):
        bm_name = bm.get('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}name')
        if bm_name and not bm_name.startswith('_'): heading_id = bm_name

    is_hr = bool(p._p.xpath('.//w:pBdr')) or p.text.strip() in ("---", "***", "___")

    for child in p._p:
        tag = child.tag.split("}")[-1]
        if tag == "hyperlink":
            anchor = child.get("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}anchor")
            rid = child.get("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id")
            txt = child.xpath("string(.)").strip()
            if not txt: continue
            if anchor:
                chunks.append(f'<a href="#{html.escape(anchor)}" class="doc-link doc-anchor-link" data-anchor="{html.escape(anchor)}">{html.escape(txt)}</a>')
            else:
                target = part.rels[rid].target_ref if (rid and rid in part.rels) else "#"
                anc = _resolve_anchor(target)
                if anc:
                    chunks.append(f'<a href="#{html.escape(anc)}" class="doc-link doc-anchor-link" data-anchor="{html.escape(anc)}">{html.escape(txt)}</a>')
                else:
                    chunks.append(f'<a href="{html.escape(target)}" target="_blank" rel="noopener" class="doc-link">{html.escape(txt)}</a>')
        elif tag == "r":
            blips = child.xpath('.//a:blip')
            if blips:
                for blip in blips:
                    rId = blip.get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed')
                    if rId and rId in part.rels:
                        try:
                            img_p = part.rels[rId].target_part
                            b64 = base64.b64encode(img_p.blob).decode('ascii')
                            ct = getattr(img_p, 'content_type', 'image/png')
                            chunks.append(f'<div class="doc-img-wrap"><img src="data:{ct};base64,{b64}" class="doc-embedded-img" alt="图表" style="max-width:100%;border-radius:4px;display:block;margin:12px auto;" /></div>')
                        except Exception: pass
            t = child.xpath("string(.)")
            if t:
                # 防御性脱壳：剥离裸露的前端 HTML 标签源码
                clean_t = re.sub(r'</?(?:div|span|p|section|article|header|footer)[^>]*>', '', t, flags=re.IGNORECASE)
                esc_t = html.escape(clean_t)
                is_c = bool(child.xpath('.//w:shd[@w:fill="F1F5F9"]') or child.xpath('.//w:rFonts[@w:ascii="Consolas"]'))
                if is_c: chunks.append(f'<code class="doc-inline-code">{esc_t}</code>')
                elif child.xpath(".//w:b"): chunks.append(f"<strong>{esc_t}</strong>")
                else: chunks.append(esc_t)

    res = "".join(chunks).strip() or html.escape(re.sub(r'</?(?:div|span|p)[^>]*>', '', p.text.strip()))

    def _md_link_sub(m):
        t, u = m.group(1), m.group(2)
        if u.startswith("#"): return f'<a href="{html.escape(u)}" class="doc-link doc-anchor-link" data-anchor="{html.escape(u[1:])}">{html.escape(t)}</a>'
        anc = _resolve_anchor(u)
        if anc: return f'<a href="#{html.escape(anc)}" class="doc-link doc-anchor-link" data-anchor="{html.escape(anc)}">{html.escape(t)}</a>'
        return f'<a href="{html.escape(u)}" target="_blank" rel="noopener" class="doc-link">{html.escape(t)}</a>'

    res = re.sub(r'\[([^\]]+)\]\(([^)]+)\)', _md_link_sub, res)
    return re.sub(r'`([^`]+)`', r'<code class="doc-inline-code">\1</code>', res), heading_id, is_hr


def _render_cell_content(cell: Any) -> str:
    """渲染表格单元格富文本并保留链接与段落换行"""
    parts = []
    for p in cell.paragraphs:
        h, _, _ = _render_paragraph_html(p)
        if h: parts.append(h)
    return "<br>".join(parts) if parts else html.escape(cell.text.strip())


def _render_single_page(items: List[Tuple[str, Any]], page_num: int, total_pages: int) -> str:
    """渲染单个标准的 A4 纸张页面（支持多语并排扉页、并排网格与正文流）"""
    if page_num == 1:
        tbl_items = [it[1] for it in items if it[0] == "tbl"]
        poly_tbl_info = next((t for t in tbl_items if len(t[0].columns) >= 2), None)
        if poly_tbl_info:
            tbl_obj = poly_tbl_info[0]
            p_texts = [it[1].text.strip() for it in items if it[0] == "p" and it[1].text.strip()]
            brand = html.escape(p_texts[0] if p_texts else "ILLACME DIGITAL BINDERY HUB")
            meta = html.escape("\n".join(p_texts[1:])).replace("\n", "<br>") if len(p_texts) > 1 else ""
            cols_html = []
            for cell in tbl_obj.rows[0].cells:
                raw_lines = [ln.strip() for ln in cell.text.splitlines() if ln.strip()]
                b_badge = html.escape(raw_lines[0]) if raw_lines else ""
                b_title = html.escape(raw_lines[1]) if len(raw_lines) > 1 else (b_badge or "数字典籍")
                b_desc = html.escape(" ".join(raw_lines[2:])) if len(raw_lines) > 2 else ""
                col_badge = f'<span class="doc-cover-badge">{b_badge}</span>' if b_badge else ""
                col_desc = f'<p class="doc-cover-col-desc">{b_desc}</p>' if b_desc else ""
                cols_html.append(f'<div class="doc-cover-col">{col_badge}<h1 class="doc-cover-col-title">{b_title}</h1>{col_desc}</div>')
            return f"""<article class="doc-page doc-cover-page" id="page-{page_num}">
  <div class="doc-cover-inner"><div class="doc-cover-brand">{brand}</div><div class="doc-cover-poly-grid">{"".join(cols_html)}</div><div class="doc-cover-divider"></div>{f'<div class="doc-cover-meta">{meta}</div>' if meta else ""}</div>
  <div class="doc-page-footer">典籍印本 · 双语扉页</div>
</article>"""
        else:
            p_texts = [it[1].text.strip() for it in items if it[0] == "p" and it[1].text.strip()]
            title = html.escape(p_texts[0] if p_texts else "典籍审校印本")
            desc = html.escape(p_texts[1]) if len(p_texts) > 1 else ""
            meta = html.escape("\n".join(p_texts[2:])).replace("\n", "<br>") if len(p_texts) > 2 else ""
            return f"""<article class="doc-page doc-cover-page" id="page-{page_num}">
  <div class="doc-cover-inner"><div class="doc-cover-brand">ILLACME DIGITAL BINDERY HUB</div><h1 class="doc-cover-title">{title}</h1>{f'<p class="doc-cover-desc">{desc}</p>' if desc else ""}<div class="doc-cover-divider"></div>{f'<div class="doc-cover-meta">{meta}</div>' if meta else ""}</div>
  <div class="doc-page-footer">典籍印本 · 扉页</div>
</article>"""

    parts = []
    for it_type, obj in items:
        if it_type == "p":
            raw_t = obj.text.strip()
            s = (obj.style.name if obj.style else "").lower()
            inner, h_id, is_hr = _render_paragraph_html(obj)
            if is_hr: parts.append('<hr class="doc-hr">'); continue
            if "<img" in inner or "doc-img-wrap" in inner: parts.append(inner); continue
            if not raw_t: continue

            m_ch = re.match(r"^第\s*(\d+)\s*章", raw_t)
            id_attr = f' id="{h_id or f"ch_{m_ch.group(1)}"}"' if (h_id or m_ch) else ""

            if "heading 1" in s or s.startswith("heading 1"): parts.append(f'<h2 class="doc-h1"{id_attr}>{inner}</h2>')
            elif "heading 2" in s or s.startswith("heading 2"): parts.append(f'<h3 class="doc-h2"{id_attr}>{inner}</h3>')
            elif "heading 3" in s or s.startswith("heading 3"): parts.append(f'<h4 class="doc-h3"{id_attr}>{inner}</h4>')
            elif raw_t.startswith(("·", "•")): parts.append(f'<div class="doc-list-item doc-sub-item">{inner}</div>')
            elif bool(re.match(r"^\d+\.\s+", raw_t)) or "list" in s or raw_t.startswith(("-", "*")):
                m_num = re.match(r"^(\d+)\.\s+", raw_t)
                target_anc = f"ch_{m_num.group(1)}" if m_num else ""
                click_wrap = f'<a href="#{target_anc}" class="doc-toc-link" data-anchor="{target_anc}">{inner}</a>' if target_anc else inner
                parts.append(f'<div class="doc-list-item doc-main-item">{click_wrap}</div>')
            elif "caption" in s: parts.append(f'<p class="doc-caption">{inner}</p>')
            elif raw_t.startswith(">"): parts.append(f'<blockquote class="doc-quote">{inner.lstrip("> ")}</blockquote>')
            else: parts.append(f'<p class="doc-para">{inner}</p>')
        elif it_type == "tbl":
            tbl_obj, s_r, e_r = obj
            if len(tbl_obj.rows) == 1 and len(tbl_obj.columns) == 1:
                cps = tbl_obj.rows[0].cells[0].paragraphs
                lt = cps[0].text.strip() if len(cps) > 1 else ""
                cb = "\n".join(p.text for p in (cps[1:] if len(cps) > 1 else cps))
                esc_c, esc_l = html.escape(cb.strip()), html.escape(lt or "CODE")
                if lt.upper() in ("MERMAID", "FLOWCHART"):
                    parts.append(
                        f'<div class="doc-code-block doc-mermaid-block" data-raw-code="{esc_c}"><div class="doc-code-header">'
                        f'<span class="doc-code-lang">📊 MERMAID 架构图</span><div class="doc-mermaid-tabs">'
                        f'<button class="doc-mermaid-tab active" onclick="switchMermaidView(this,\'diagram\')">📊 流程图</button>'
                        f'<button class="doc-mermaid-tab" onclick="switchMermaidView(this,\'code\')">💻 源代码</button>'
                        f'<button class="doc-code-copy" onclick="copyMermaidCode(this)">复制</button></div></div>'
                        f'<div class="doc-mermaid-diagram-wrap"><div class="mermaid">{esc_c}</div></div>'
                        f'<div class="doc-mermaid-code-wrap" style="display:none;"><pre><code>{esc_c}</code></pre></div></div>'
                    )
                else:
                    parts.append(f'<div class="doc-code-block"><div class="doc-code-header"><span class="doc-code-lang">{esc_l}</span><button class="doc-code-copy" onclick="navigator.clipboard.writeText(this.closest(\'.doc-code-block\').querySelector(\'code\').innerText).then(()=>{{this.textContent=\'已复制\';setTimeout(()=>this.textContent=\'复制\',1500)}})">复制</button></div><pre><code>{esc_c}</code></pre></div>')
                continue

            r0_txt = " ".join(c.text for c in tbl_obj.rows[0].cells)
            is_poly = any(k in r0_txt for k in ("[ZH]", "[EN]", "[JA]", "[FR]", "[DE]", "目录", "Contents", "版权", "Colophon"))
            tbl_tag = '<table class="doc-polyglot-grid">' if is_poly else '<table>'

            table_rows = []
            # 如果是切片后续，且原表有表头，自动追加表头保证跨页对齐
            if s_r > 0 and is_poly and len(tbl_obj.rows) > 0:
                h_cells = "".join([f'<th>{_render_cell_content(c)}</th>' for c in tbl_obj.rows[0].cells])
                table_rows.append(f'<tr>{h_cells}</tr>')

            for r_idx in range(s_r, min(e_r, len(tbl_obj.rows))):
                r = tbl_obj.rows[r_idx]
                tag = 'th' if (r_idx == 0 and s_r == 0) else 'td'
                cells_rendered = [f'<{tag}>{_render_cell_content(c)[8:-9] if (tag == "th" and _render_cell_content(c).startswith("<strong>") and _render_cell_content(c).endswith("</strong>")) else _render_cell_content(c)}</{tag}>' for c in r.cells]
                table_rows.append(f'<tr>{"".join(cells_rendered)}</tr>')
            if table_rows: parts.append(f'<div class="doc-table-wrap">{tbl_tag}{"".join(table_rows)}</table></div>')

    content_body = "\n    ".join(parts)
    return f"""<article class="doc-page" id="page-{page_num}">
  <div class="doc-page-content">{content_body}</div>
  <div class="doc-page-footer">第 {page_num} 页 · 共 {total_pages} 页</div>
</article>"""

